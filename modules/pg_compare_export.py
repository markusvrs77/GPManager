# -*- coding: utf-8 -*-
"""
Выгрузка результатов сравнения баз PostgreSQL (задача pg_compare) в Excel.

Только чтение SQLite: результаты (pg_compare.get_results), несовпавшие
диапазоны (pg_compare.get_mismatched_ranges) и имена подключений —
без паролей и прочих реквизитов.
"""

import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import modules.pg_compare as pg_compare
from modules.connections import get_connection_by_id

# подписи — как в UI (static/js/pg_compare.js)
STATUS_RU = {
    "same": "Совпадает",
    "differs": "Отличается",
    "no_dest": "Нет в приёмнике",
    "no_source": "Нет в источнике",
    "structure_diff": "Структура отличается",
    "duplicate_keys": "Дубликаты ключа",
    "error": "Ошибка",
    "cancelled": "Отменено",
}
KEY_SOURCE_RU = {"pk": "PK", "unique_index": "уник. индекс",
                 "sync_keys": "сохранённый ключ"}
JOB_STATUS_RU = {
    "queued": "в очереди", "pending": "в очереди", "running": "идёт",
    "stopping": "останавливается", "done": "завершено", "failed": "ошибка",
    "cancelled": "остановлено", "interrupted": "прервано",
}

# фильтр → статусы (None — все)
FILTERS = {
    "all": None,
    "needs": ("differs", "no_dest"),
    "same": ("same",),
    "problems": ("structure_diff", "duplicate_keys", "error", "no_source",
                 "cancelled"),
}
FILTER_RU = {"all": "все", "needs": "нужно выровнять", "same": "совпадают",
             "problems": "проблемы"}

RESULT_HEADER = ["Схема", "Таблица", "Статус", "Ключ", "Источник ключа",
                 "Строк в источнике", "Строк в приёмнике", "Добавить",
                 "Изменить", "Удалить", "Всего отличий", "Диапазоны",
                 "Сообщение", "Время сравнения", "Приёмник"]
RESULT_NUMERIC = (6, 7, 8, 9, 10, 11)

RANGE_HEADER = ["Схема", "Таблица", "Колонка", "Режим", "От", "До",
                "NULL-диапазон", "Строк в источнике", "Строк в приёмнике",
                "Добавить", "Изменить", "Удалить"]
RANGE_NUMERIC = (8, 9, 10, 11, 12)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def total_diff(row):
    """Добавить + изменить + удалить; пустое считается нулём."""
    return sum(int(row.get(k) or 0)
               for k in ("to_insert", "to_update", "to_delete"))


def dest_name(row):
    """Таблица приёмника: цель карты targets или то же schema.table."""
    return row.get("target") or "%s.%s" % (row.get("schema"),
                                           row.get("table"))


def select_rows(results, filter_name, query):
    """Строки под фильтром и поиском по schema.table источника или цели
    (без учёта регистра), по «всего отличий» по убыванию."""
    statuses = FILTERS[filter_name]
    needle = (query or "").strip().lower()

    rows = [r for r in results
            if (statuses is None or r.get("status") in statuses)
            and (not needle
                 or needle in ("%s.%s" % (r.get("schema"),
                                          r.get("table"))).lower()
                 or needle in str(r.get("target") or "").lower())]
    rows.sort(key=lambda r: (-total_diff(r), str(r.get("schema")),
                             str(r.get("table"))))
    return rows


def connection_name(connection_id):
    """Имя подключения — и ничего больше из его реквизитов."""
    try:
        conn = get_connection_by_id(int(connection_id))
    except (TypeError, ValueError):
        conn = None
    if not conn:
        return "#%s" % connection_id
    return "%s (#%s)" % (conn.get("name") or "", conn.get("id"))


def _text(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return "да" if value else "нет"
    if isinstance(value, (int, float)):
        return value
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False)
    return ILLEGAL_CHARACTERS_RE.sub("", value)


def _append(ws, values):
    ws.append([_text(v) for v in values])
    # строка, начинающаяся с «=», — текст, а не формула
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def _style_sheet(ws, numeric_cols, max_width=60):
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1 or ws.max_column > 1:
        ws.auto_filter.ref = ws.dimensions

    for col in numeric_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = "#,##0"

    for idx, column in enumerate(ws.columns, start=1):
        width = max(len("" if c.value is None else str(c.value))
                    for c in column)
        ws.column_dimensions[get_column_letter(idx)].width = \
            max(10, min(width + 3, max_width))


def _ranges_text(chunked):
    if not chunked:
        return None
    return "%s / %s / %s" % (chunked.get("checked"), chunked.get("total"),
                             chunked.get("mismatched"))


def _results_sheet(ws, rows):
    ws.title = "Результаты"
    ws.append(RESULT_HEADER)
    for r in rows:
        _append(ws, [
            r.get("schema"), r.get("table"),
            STATUS_RU.get(r.get("status"), r.get("status")),
            ", ".join(r.get("key_columns") or []) or None,
            KEY_SOURCE_RU.get(r.get("key_source"), r.get("key_source")),
            r.get("src_rows"), r.get("dst_rows"),
            r.get("to_insert"), r.get("to_update"), r.get("to_delete"),
            total_diff(r), _ranges_text(r.get("chunked")),
            r.get("message"), r.get("compared_at"), dest_name(r),
        ])
    _style_sheet(ws, RESULT_NUMERIC)


def _ranges_sheet(wb, job_id, rows):
    found = []
    for r in rows:
        if not r.get("chunked"):
            continue
        for rng in pg_compare.get_mismatched_ranges(job_id, r["schema"],
                                                    r["table"]):
            found.append((r, rng))
    if not found:
        return

    ws = wb.create_sheet("Несовпавшие диапазоны")
    ws.append(RANGE_HEADER)
    for r, rng in found:
        _append(ws, [
            r["schema"], r["table"], rng.get("column"),
            "корзина" if rng.get("mode") == "bucket" else "диапазон",
            rng.get("lo"), rng.get("hi"), bool(rng.get("is_null")),
            rng.get("src_rows"), rng.get("dst_rows"), rng.get("to_insert"),
            rng.get("to_update"), rng.get("to_delete"),
        ])
    _style_sheet(ws, RANGE_NUMERIC)


def _summary_sheet(wb, job, config, results, filter_name, query, now):
    ws = wb.create_sheet("Сводка")
    ws.append(["Параметр", "Значение"])
    counts = {}
    for r in results:
        counts[r.get("status")] = counts.get(r.get("status"), 0) + 1

    lines = [
        ("Сравнение", "#%s" % job["id"]),
        ("Источник", connection_name(config.get("source_connection_id"))),
        ("Приёмник", connection_name(config.get("dest_connection_id"))),
        ("Состояние задачи",
         JOB_STATUS_RU.get(job.get("status"), job.get("status"))),
        ("Начато", job.get("started_at") or job.get("created_at")),
        ("Завершено", job.get("finished_at")),
        ("Выгружено", now.strftime("%Y-%m-%d %H:%M:%S")),
        ("Фильтр", FILTER_RU[filter_name]),
        ("Поиск", query or None),
        ("Таблиц всего", len(results)),
    ]
    for status in pg_compare.STATUSES:
        lines.append((STATUS_RU.get(status, status), counts.get(status, 0)))
    for line in lines:
        _append(ws, line)
    _style_sheet(ws, (), max_width=80)


def _job_config(job):
    try:
        cfg = json.loads(job.get("config_json") or "{}")
    except (TypeError, ValueError):
        cfg = {}
    return cfg if isinstance(cfg, dict) else {}


def build_export(job, filter_name="all", query="", now=None):
    """(BytesIO с книгой, имя файла). filter_name — ключ FILTERS."""
    if filter_name not in FILTERS:
        raise ValueError("Неизвестный фильтр: %s" % filter_name)
    now = now or datetime.now()

    results = pg_compare.get_results(job["id"])
    rows = select_rows(results, filter_name, query)

    wb = Workbook()
    _results_sheet(wb.active, rows)
    _ranges_sheet(wb, job["id"], rows)
    _summary_sheet(wb, job, _job_config(job), results, filter_name, query,
                   now)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out, "pg_compare_%s_%s.xlsx" % (job["id"],
                                           now.strftime("%Y%m%d_%H%M"))
