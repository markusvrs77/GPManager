# -*- coding: utf-8 -*-
"""
Загрузка в другую таблицу приёмника (карта targets) в Postgres Toolkit:
перенос copy_pipe, сравнение pg_compare, загрузка разницы pg_diff_load,
результаты и выгрузка в Excel. Без живой базы — фейковые соединения.
"""

import io
import json
import threading

import pytest
from openpyxl import load_workbook

import modules.pg_compare as cmp
import modules.pg_diff_load as pdl
import modules.pg_ranges as pr
import modules.pg_sync_common as common
import modules.sync_transport as st
from job_manager import create_job, get_job, get_job_items, mark_job_done
from tests.pg_fakes import FakeConn

PG = {"host": "h", "port": 5432, "username": "u", "database_name": "cash",
      "db_type": "postgres"}


def _catalog(namespaces, relations, inherits=()):
    """Фейк каталога одной стороны (как в test_pg_compare)."""
    return FakeConn(responses=[
        ("pg_inherits", list(inherits)),
        ("c.relkind IN", list(relations)),
        ("FROM pg_namespace", [(n,) for n in namespaces]),
    ])


# ------------------------------------------------------------ expand_selection

def test_mapped_table_is_in_dest_by_its_target_name():
    src = _catalog(["sales"], [("sales", "orders"), ("sales", "items")])
    dst = _catalog(["sales", "arch"], [("arch", "orders_copy"),
                                       ("sales", "items")])

    result = cmp.expand_selection(src, dst, ["sales"], [],
                                  targets={"sales.orders": "arch.orders_copy"})

    rows = {r["table"]: r for r in result}
    assert rows["orders"]["in_dst"] is True
    assert rows["orders"]["target"] == "arch.orders_copy"
    assert rows["items"]["in_dst"] is True and "target" not in rows["items"]
    # каталог приёмника читается и по схеме цели
    assert any(params and "arch" in params[0]
               for _text, params in dst.executed)


def test_missing_target_is_no_dest_even_if_same_name_exists():
    src = _catalog(["sales"], [("sales", "orders")])
    dst = _catalog(["sales"], [("sales", "orders")])

    result = cmp.expand_selection(
        src, dst, [], [{"schema": "sales", "table": "orders"}],
        targets={"sales.orders": "orders_new"})

    assert result == [{"schema": "sales", "table": "orders", "in_src": True,
                       "in_dst": False, "target": "sales.orders_new"}]


def test_target_is_not_listed_as_only_in_dest():
    src = _catalog(["sales"], [("sales", "x")])
    dst = _catalog(["sales"], [("sales", "y"), ("sales", "legacy")])

    result = cmp.expand_selection(src, dst, ["sales"], [],
                                  targets={"sales.x": "sales.y"})

    assert sorted((r["table"], r["in_src"], r["in_dst"]) for r in result) == [
        ("legacy", False, True), ("x", True, True)]


@pytest.mark.parametrize("targets, words", [
    ({"sales.orders": "Bad.Name"}, "строчные"),
    ({"sales.ghost": "sales.z"}, "не выбрана"),
    ({"sales.orders": "sales.items"}, "своя цель"),
    ({"sales.orders": "a.b.c"}, "schema.table"),
])
def test_bad_targets_are_rejected(targets, words):
    src = _catalog(["sales"], [("sales", "orders"), ("sales", "items")])
    dst = _catalog(["sales"], [])

    with pytest.raises(ValueError) as err:
        cmp.expand_selection(src, dst, ["sales"], [], targets=targets)

    assert words in str(err.value)


# ------------------------------------------------------------ compare_table

COLS = [("id", "integer"), ("name", "text")]


def _sides():
    src = FakeConn(responses=[("a.attname", list(COLS))],
                   copy_out=b"1\ta\n")
    dst = FakeConn(responses=[
        ("indisunique", []),
        ("a.attname", list(COLS)),
        ('FROM "pgcmp_src" GROUP BY', [(0,)]),
        ("HAVING count(*) > 1", [(0,)]),
        ("AS to_insert", [(1, 0, 0, 0)]),
    ])
    return src, dst


def test_compare_reads_source_by_its_name_and_dest_by_the_target():
    src, dst = _sides()

    result = cmp.compare_table(src, dst, "sales", "orders", ["id"],
                               dst=("arch", "orders_copy"))

    assert result["status"] == "same"
    assert 'FROM "sales"."orders" AS "t"' in src.copies[0]
    dst_sql = dst.sql_text()
    assert '"arch"."orders_copy"' in dst_sql
    assert '"sales"."orders"' not in dst_sql
    # каталог приёмника — по цели: колонки и уникальность ключа
    dst_params = [p for _t, p in dst.executed if p]
    assert ("arch", "orders_copy") in [tuple(p[:2]) for p in dst_params]
    assert ("sales", "orders") not in [tuple(p[:2]) for p in dst_params]
    # источник не меняется
    assert not any(t.startswith(("INSERT", "UPDATE", "DELETE", "CREATE",
                                 "DROP", "TRUNCATE"))
                   for t, _p in src.executed)


def test_chunk_column_reads_dest_catalog_by_the_target():
    src = FakeConn(responses=[("a.attname", [("id", "integer", None, None,
                                              None)])])
    dst = FakeConn(responses=[("a.attname", [("id", "integer", None, None,
                                              None)])])

    column = pr.pick_chunk_column(src, dst, "sales", "orders", ["id"],
                                  dst=("arch", "orders_copy"))

    assert column["name"] == "id"
    assert dst.executed[0][1] == ("arch", "orders_copy")
    assert src.executed[0][1] == ("sales", "orders")


def test_runner_compares_with_the_target_and_stores_it(monkeypatch):
    calls = []

    monkeypatch.setattr(cmp, "open_pg",
                        lambda cid, readonly=False: FakeConn())
    monkeypatch.setattr(cmp, "resolve_key_candidates",
                        lambda source_id, tables: {})
    monkeypatch.setattr(cmp.pg_ranges, "should_chunk", lambda c, s, t: False)
    monkeypatch.setattr(cmp, "table_columns",
                        lambda conn, s, t: calls.append(("cols", s, t))
                        or ["id"])
    monkeypatch.setattr(common, "is_stop_requested", lambda job_id: False)

    def fake_compare(src_conn, dst_conn, schema, table, key_columns,
                     key_source=None, work_mem=None, where=None, **kw):
        calls.append(("compare", schema, table, kw.get("dst")))
        return cmp._result("same", src_rows=1, dst_rows=1, to_insert=0,
                           to_update=0, to_delete=0)

    monkeypatch.setattr(cmp, "compare_table", fake_compare)

    job_id = create_job("pg_compare", 11, {
        "source_connection_id": 11, "dest_connection_id": 12,
        "tables": [{"schema": "s", "table": "a", "in_src": True,
                    "in_dst": True, "target": "s.a_copy"},
                   {"schema": "s", "table": "b", "in_src": True,
                    "in_dst": True}],
        "targets": {"s.a": "s.a_copy"}, "parallel": 1,
        "item_action": "COMPARE"})

    cmp.run_pg_compare_job(job_id)

    assert get_job(job_id)["status"] == "done"
    assert ("compare", "s", "a", ("s", "a_copy")) in calls
    assert ("compare", "s", "b", None) in calls
    assert ("cols", "s", "a_copy") in calls
    results = {r["table"]: r for r in cmp.get_results(job_id)}
    assert results["a"]["target"] == "s.a_copy"
    assert results["b"]["target"] is None
    # строки задачи — по имени источника
    assert sorted(i["table_name"] for i in get_job_items(job_id)) == ["a",
                                                                      "b"]


def test_missing_target_is_no_dest_with_its_name():
    plan = cmp._compare_one(FakeConn(), FakeConn(), "s", "a",
                            {"in_src": True, "in_dst": False}, [],
                            dst=("s", "a_copy"))

    assert plan["row"]["status"] == "no_dest"
    assert "s.a_copy" in plan["row"]["message"]


# ------------------------------------------------------------ маршрут сравнения

@pytest.fixture
def pg_routes(monkeypatch):
    """31, 32 — Postgres; каталог — фейки; раннеры не стартуют."""
    kinds = {31: dict(PG, id=31, name="prod"), 32: dict(PG, id=32, name="test")}
    started = []
    sides = {"src": _catalog(["sales"], [("sales", "orders"),
                                          ("sales", "items")]),
             "dst": _catalog(["sales"], [("sales", "orders_copy"),
                                         ("sales", "items")])}

    monkeypatch.setattr(common, "get_connection_by_id",
                        lambda cid: kinds.get(int(cid)))
    monkeypatch.setattr(common, "open_pg", lambda cid, readonly=False:
                        sides["src"] if readonly else sides["dst"])

    class NoThread(object):
        def __init__(self, target=None, args=(), daemon=None, **_kw):
            self.target, self.args = target, args

        def start(self):
            started.append((self.target, self.args))

    monkeypatch.setattr(threading, "Thread", NoThread)
    return started


def _config(job_id):
    return json.loads(get_job(job_id)["config_json"])


def test_compare_start_stores_normalized_targets(client, pg_routes):
    response = client.post("/api/pg/compare/start", json={
        "source_connection_id": 31, "dest_connection_id": 32,
        "schemas": ["sales"], "tables": [],
        "targets": {"sales.orders": "orders_copy", "sales.items": ""}})
    body = response.get_json()

    assert response.status_code == 200, body
    config = _config(body["job_id"])
    assert config["targets"] == {"sales.orders": "sales.orders_copy"}
    tables = {t["table"]: t for t in config["tables"]}
    assert tables["orders"]["in_dst"] is True
    assert tables["orders"]["target"] == "sales.orders_copy"


def test_compare_start_without_targets_keeps_old_config(client, pg_routes):
    response = client.post("/api/pg/compare/start", json={
        "source_connection_id": 31, "dest_connection_id": 32,
        "schemas": ["sales"], "tables": []})

    assert response.status_code == 200
    assert "targets" not in _config(response.get_json()["job_id"])


@pytest.mark.parametrize("targets, words", [
    ({"sales.orders": "Orders"}, "строчные"),
    ({"sales.ghost": "x"}, "не выбрана"),
    ({"sales.orders": "sales.items"}, "своя цель"),
    ("sales.orders", "объектом"),
])
def test_compare_start_refuses_bad_targets(client, pg_routes, targets, words):
    response = client.post("/api/pg/compare/start", json={
        "source_connection_id": 31, "dest_connection_id": 32,
        "schemas": ["sales"], "tables": [], "targets": targets})

    assert response.status_code == 400
    assert words in response.get_json()["message"]
    assert pg_routes == []


# ------------------------------------------------------------ маршрут загрузки

@pytest.fixture
def mapped_compare():
    """Законченное сравнение 31 → 32: orders сравнивали с orders_copy."""
    job_id = create_job("pg_compare", 31, {
        "source_connection_id": 31, "dest_connection_id": 32,
        "tables": [{"schema": "sales", "table": "orders",
                    "target": "sales.orders_copy"},
                   {"schema": "sales", "table": "items"}],
        "targets": {"sales.orders": "sales.orders_copy"},
        "item_action": "COMPARE"})
    cmp.save_result(job_id, {"schema": "sales", "table": "orders",
                             "status": "differs", "key_columns": ["id"],
                             "key_source": "pk", "to_insert": 1,
                             "to_update": 0, "to_delete": 0,
                             "target": "sales.orders_copy"})
    cmp.save_result(job_id, {"schema": "sales", "table": "items",
                             "status": "same"})
    mark_job_done(job_id)
    return job_id


def _load(client, **body):
    payload = {"source_connection_id": 31, "dest_connection_id": 32,
               "delete_missing": False,
               "tables": [{"schema": "sales", "table": "orders",
                           "action": "diff"}]}
    payload.update(body)
    return client.post("/api/pg/diff-load/start", json=payload)


def test_diff_load_takes_targets_from_the_compare(client, pg_routes,
                                                  mapped_compare):
    response = _load(client, compare_job_id=mapped_compare)
    body = response.get_json()

    assert response.status_code == 200, body
    config = _config(body["job_id"])
    assert config["targets"] == {"sales.orders": "sales.orders_copy"}
    assert config["tables"][0]["target"] == "sales.orders_copy"


def test_diff_load_accepts_the_same_targets(client, pg_routes,
                                            mapped_compare):
    response = _load(client, compare_job_id=mapped_compare,
                     targets={"sales.orders": "orders_copy",
                              "sales.items": ""})

    assert response.status_code == 200, response.get_json()


@pytest.mark.parametrize("targets", [
    {"sales.orders": "sales.other"},
    {"sales.orders": ""},
])
def test_diff_load_refuses_targets_other_than_the_compare(
        client, pg_routes, mapped_compare, targets):
    response = _load(client, compare_job_id=mapped_compare, targets=targets)

    assert response.status_code == 400
    assert "отличается от сравнения" in response.get_json()["message"]
    assert pg_routes == []


def test_diff_load_refuses_new_targets_for_an_unmapped_compare(
        client, pg_routes, mapped_compare):
    response = _load(client, compare_job_id=mapped_compare,
                     tables=[{"schema": "sales", "table": "items",
                              "action": "diff"}],
                     targets={"sales.items": "items_2"})

    assert response.status_code == 400
    assert pg_routes == []


def test_full_load_without_compare_accepts_request_targets(client,
                                                           pg_routes):
    response = _load(client, tables=[
        {"schema": "sales", "table": "orders", "action": "full"}],
        targets={"sales.orders": "orders_copy"})
    body = response.get_json()

    assert response.status_code == 200, body
    config = _config(body["job_id"])
    assert config["targets"] == {"sales.orders": "sales.orders_copy"}
    assert config["tables"][0]["in_dst"] is True
    assert config["tables"][0]["target"] == "sales.orders_copy"


def test_full_load_without_compare_refuses_bad_targets(client, pg_routes):
    response = _load(client, tables=[
        {"schema": "sales", "table": "orders", "action": "full"}],
        targets={"sales.orders": "bad name"})

    assert response.status_code == 400
    assert pg_routes == []


# ------------------------------------------------------------ загрузка разницы

@pytest.fixture
def typed(monkeypatch):
    """Типы колонок по сторонам и по имени таблицы."""
    seen = []

    def fake_types(conn, schema, table):
        seen.append((conn.side, schema, table))
        return {"id": "integer", "name": "text"}

    monkeypatch.setattr(pdl, "table_column_types", fake_types)
    return seen


def _pair(dst_responses=()):
    src = FakeConn(copy_out=b"1\ta\n")
    src.side = "src"
    dst = FakeConn(responses=list(dst_responses))
    dst.side = "dst"
    return src, dst


KEYED = [("GROUP BY", []), ("IS NULL", [(0,)]), ("DELETE FROM", [(1,)]),
         ("UPDATE", [(1,)]), ("INSERT INTO", [(1,)])]
TARGET = '"arch"."orders_copy"'


def test_diff_writes_only_the_target_and_reads_the_source(typed):
    src, dst = _pair(KEYED)

    out = pdl.load_diff(src, dst, "sales", "orders", ["id"], True, "stg_7_1",
                        dst=("arch", "orders_copy"))

    assert out == {"insert": 1, "update": 1, "delete": 1}
    assert ("src", "sales", "orders") in typed
    assert ("dst", "arch", "orders_copy") in typed
    texts = [t for t, _p in dst.executed]
    for verb in ("CREATE UNLOGGED TABLE", "DELETE FROM", "UPDATE",
                 "INSERT INTO"):
        text = next(t for t in texts if t.startswith(verb))
        assert TARGET in text, text
        assert '"sales"."orders"' not in text, text
    # источник только читается и по своему имени
    assert src.copies == ['COPY (SELECT "id", "name" FROM "sales"."orders") '
                          'TO STDOUT']
    # флаги колонок — по цели
    flags = [p for t, p in dst.executed if "attgenerated" in t]
    assert flags == [("arch", "orders_copy")]


def test_full_load_truncates_and_fills_the_target(typed):
    src, dst = _pair()

    pdl.load_full(src, dst, "sales", "orders", truncate=True,
                  dst=("arch", "orders_copy"))

    texts = [t for t, _p in dst.executed]
    assert 'TRUNCATE TABLE "arch"."orders_copy"' in texts
    assert dst.copies == ['COPY "arch"."orders_copy" ("id", "name") '
                          'FROM STDIN']
    assert src.copies == ['COPY (SELECT "id", "name" FROM "sales"."orders") '
                          'TO STDOUT']


def test_create_builds_the_target_from_the_source_definition():
    src = FakeConn(responses=[
        ("format_type", [("id", "integer", True, "", "", None, "r", None,
                          None)]),
        ("contype = 'p'", [("id",)]),
    ])
    dst = FakeConn(responses=[("FROM pg_namespace", [(1,)])])

    pdl.create_table_from_source(src, dst, "sales", "orders",
                                 dst=("arch", "orders_copy"))

    assert src.executed[1][1] == ("sales", "orders")
    texts = [t for t, _p in dst.executed]
    assert any(t.startswith('CREATE TABLE "arch"."orders_copy" ("id" integer '
                            'NOT NULL, PRIMARY KEY ("id"))') for t in texts)
    assert dst.executed[0][1] == ("arch",)


def test_sequences_and_existence_use_the_target():
    dst = FakeConn(responses=[("pg_get_serial_sequence",
                               [("id", "arch.orders_copy_id_seq")])])

    pdl.sync_sequences(dst, "arch", "orders_copy")

    texts = [(t, p) for t, p in dst.executed]
    assert texts[0][1] == ("arch", "orders_copy")
    assert 'FROM "arch"."orders_copy"' in texts[1][0]


def test_runner_loads_each_table_into_its_target(monkeypatch):
    state = {"diff": [], "full": [], "create": [], "exists": [], "seq": []}

    monkeypatch.setattr(pdl, "open_pg", lambda cid, readonly=False:
                        FakeConn())
    monkeypatch.setattr(common, "is_stop_requested", lambda job_id: False)

    def fake_diff(src_conn, dst_conn, schema, table, key_columns, delete_missing,
                  stage_name, **kw):
        state["diff"].append((schema, table, kw.get("dst")))
        return {"insert": 1, "update": 0, "delete": 0}

    def fake_full(src_conn, dst_conn, schema, table, truncate, require_empty=False,
                  **kw):
        state["full"].append((schema, table, kw.get("dst")))
        return {"rows": 2}

    def fake_create(src_conn, dst_conn, schema, table, **kw):
        state["create"].append((schema, table, kw.get("dst")))
        return {"partitioned": False, "virtual_as_stored": []}

    monkeypatch.setattr(pdl, "load_diff", fake_diff)
    monkeypatch.setattr(pdl, "load_full", fake_full)
    monkeypatch.setattr(pdl, "create_table_from_source", fake_create)
    monkeypatch.setattr(pdl, "dest_table_exists", lambda conn, s, t:
                        state["exists"].append((s, t)) or False)
    monkeypatch.setattr(pdl, "sync_sequences", lambda conn, s, t:
                        state["seq"].append((s, t)) or [])

    job_id = create_job("pg_diff_load", 11, {
        "source_connection_id": 11, "dest_connection_id": 12,
        "compare_job_id": None, "delete_missing": False,
        "tables": [{"schema": "s", "table": "a", "action": "diff",
                    "key_columns": ["id"], "in_dst": True},
                   {"schema": "s", "table": "b", "action": "full",
                    "key_columns": [], "in_dst": True},
                   {"schema": "s", "table": "c", "action": "create",
                    "key_columns": [], "in_dst": False},
                   {"schema": "s", "table": "d", "action": "full",
                    "key_columns": [], "in_dst": True}],
        "targets": {"s.a": "s.a2", "s.b": "t.b2", "s.c": "s.c2"},
        "expected": []})

    pdl.run_pg_diff_load_job(job_id)

    assert get_job(job_id)["status"] == "done"
    assert state["diff"] == [("s", "a", ("s", "a2"))]
    assert ("s", "b", ("t", "b2")) in state["full"]
    # create: создаётся и заливается цель
    assert state["exists"] == [("s", "c2")]
    assert state["create"] == [("s", "c", ("s", "c2"))]
    assert ("s", "c", ("s", "c2")) in state["full"]
    # без карты — прежний вызов, без dst
    assert ("s", "d", None) in state["full"]
    assert state["seq"] == [("s", "a2"), ("t", "b2"), ("s", "c2"),
                            ("s", "d")]
    assert {i["table_name"]: i["status"] for i in get_job_items(job_id)} == {
        "a": "done", "b": "done", "c": "done", "d": "done"}


# ------------------------------------------------------------ copy_pipe

def test_copy_pipe_sql_with_explicit_columns():
    out, into = st.build_copy_pipe_sql("sales", "orders", "arch", "copy",
                                       ["id", "na\"me"])

    assert out == ('COPY (SELECT "id", "na""me" FROM "sales"."orders") '
                   'TO STDOUT')
    assert into == 'COPY "arch"."copy" ("id", "na""me") FROM STDIN'


def test_copy_pipe_sql_without_columns_is_unchanged():
    assert st.build_copy_pipe_sql("s", "t", "s", "t") == (
        'COPY (SELECT * FROM "s"."t") TO STDOUT', 'COPY "s"."t" FROM STDIN')


def test_ensure_dest_table_creates_the_target_from_the_source():
    src = FakeConn(responses=[("a.attname", [("id", "integer"),
                                             ("name", "text")])])
    dst = FakeConn(responses=[("FROM pg_class", [])])

    created = st.ensure_dest_table(src, dst, "sales", "orders", False,
                                   dst_schema="arch", dst_table="copy")

    assert created is True
    assert dst.executed[0][1] == ("arch", "copy")
    assert src.executed[0][1] == ("sales", "orders")
    texts = [t for t, _p in dst.executed]
    assert 'CREATE SCHEMA IF NOT EXISTS "arch"' in texts
    assert any(t.startswith('CREATE TABLE "arch"."copy"') for t in texts)
    assert dst.commits == 1


def test_mapped_columns_follow_source_order_and_must_exist_in_target():
    src = FakeConn(responses=[("a.attname", [("id", "integer"),
                                             ("name", "text")])])
    dst = FakeConn(responses=[("a.attname", [("name", "text"),
                                             ("id", "integer"),
                                             ("extra", "text")])])

    assert st.mapped_copy_columns(src, dst, "s", "t", "a", "b") == ["id",
                                                                    "name"]

    narrow = FakeConn(responses=[("a.attname", [("id", "integer")])])
    with pytest.raises(ValueError) as err:
        st.mapped_copy_columns(src, narrow, "s", "t", "a", "b")
    assert "name" in str(err.value)


def test_copy_pipe_runner_uses_target_and_explicit_columns(monkeypatch):
    calls = {"ensure": [], "copy": []}

    monkeypatch.setattr(st, "get_connection_by_id", lambda cid: PG)
    monkeypatch.setattr(st, "open_psycopg2_connection_by_cfg",
                        lambda cfg: FakeConn())
    monkeypatch.setattr(st, "fetch_table_sizes", lambda conn, tables: {})
    monkeypatch.setattr(st, "mapped_copy_columns",
                        lambda *a: ["id", "name"])

    def fake_ensure(src, dst, schema, table, **kw):
        calls["ensure"].append((schema, table, kw.get("dst_schema"),
                                kw.get("dst_table")))

    def fake_copy(src, dst, s_schema, s_table, d_schema, d_table, **kw):
        calls["copy"].append((s_schema, s_table, d_schema, d_table,
                              kw.get("columns")))

    monkeypatch.setattr(st, "ensure_dest_table", fake_ensure)
    monkeypatch.setattr(st, "copy_table_pipe", fake_copy)

    job_id = create_job("copy_pipe", 11, {
        "source_connection_id": 11, "dest_connection_id": 12,
        "tables": [{"schema": "s", "table": "a"},
                   {"schema": "s", "table": "b"}],
        "targets": {"s.a": "arch.a_copy"}})

    st.run_copy_pipe_job(job_id)

    assert get_job(job_id)["status"] == "done"
    assert calls["ensure"] == [("s", "a", "arch", "a_copy"),
                               ("s", "b", None, None)]
    assert calls["copy"] == [("s", "a", "arch", "a_copy", ["id", "name"]),
                             ("s", "b", "s", "b", None)]


# ------------------------------------------------------------ Excel

def test_export_has_dest_column_and_finds_by_target(client):
    job_id = create_job("pg_compare", 91, {
        "source_connection_id": 91, "dest_connection_id": 92,
        "tables": [], "item_action": "COMPARE"})
    cmp.save_result(job_id, {"schema": "s", "table": "a", "status": "same",
                             "target": "arch.a_copy"})
    cmp.save_result(job_id, {"schema": "s", "table": "b", "status": "same"})

    response = client.get("/api/pg/compare/%d/export.xlsx" % job_id)
    ws = load_workbook(io.BytesIO(response.get_data()))["Результаты"]
    header = [c.value for c in ws[1]]
    column = header.index("Приёмник")
    rows = {r[1]: r[column] for r in ws.iter_rows(min_row=2,
                                                  values_only=True)}

    assert rows == {"a": "arch.a_copy", "b": "s.b"}

    found = client.get("/api/pg/compare/%d/export.xlsx?q=a_copy" % job_id)
    ws = load_workbook(io.BytesIO(found.get_data()))["Результаты"]
    assert [r[1] for r in ws.iter_rows(min_row=2, values_only=True)] == ["a"]


def test_results_route_carries_target(client):
    job_id = create_job("pg_compare", 93, {
        "source_connection_id": 93, "dest_connection_id": 94,
        "tables": [{"schema": "s", "table": "a"}], "item_action": "COMPARE"})
    cmp.save_result(job_id, {"schema": "s", "table": "a", "status": "same",
                             "target": "s.a_copy"})

    body = client.get("/api/pg/compare/results?job_id=%d" % job_id).get_json()

    assert body["results"][0]["target"] == "s.a_copy"


# ------------------------------------------------------------ конфиг без маршрута

# карта из сохранённого конфига (расписание, перезапуск) маршрут не видел:
# раннер проверяет её сам и валит задачу до первой таблицы
BAD_MAPS = [
    {"s.a": "s.z", "s.b": "s.z"},   # две таблицы в одну цель
    {"s.a": "s.b"},                 # цель — другая выбранная таблица
    {"s.a": "nodot!"},              # формат имени
    {"s.ghost": "s.z"},             # ключ вне задачи
    ["s.a"],                        # не объект
]


def _two_tables():
    return [{"schema": "s", "table": "a"}, {"schema": "s", "table": "b"}]


def _untouched(job_id):
    job = get_job(job_id)
    assert job["status"] == "failed"
    assert "куда грузить" in job["error_message"]
    # незакрытые строки mark_job_failed закрывает как interrupted; done нет
    assert {i["status"] for i in get_job_items(job_id)} <= {
        "queued", "pending", "interrupted"}


@pytest.mark.parametrize("targets", BAD_MAPS)
def test_copy_pipe_runner_rejects_a_bad_stored_map(monkeypatch, targets):
    touched = []
    monkeypatch.setattr(st, "get_connection_by_id", lambda cid: PG)
    monkeypatch.setattr(st, "open_psycopg2_connection_by_cfg",
                        lambda cfg: touched.append("open") or FakeConn())
    monkeypatch.setattr(st, "ensure_dest_table",
                        lambda *a, **k: touched.append("ddl"))
    monkeypatch.setattr(st, "copy_table_pipe",
                        lambda *a, **k: touched.append("copy"))

    job_id = create_job("copy_pipe", 11, {
        "source_connection_id": 11, "dest_connection_id": 12,
        "truncate": True, "tables": _two_tables(), "targets": targets})

    st.run_copy_pipe_job(job_id)

    _untouched(job_id)
    assert touched == []


@pytest.mark.parametrize("targets", BAD_MAPS)
def test_compare_runner_rejects_a_bad_stored_map(monkeypatch, targets):
    opened = []
    monkeypatch.setattr(cmp, "open_pg", lambda cid, readonly=False:
                        opened.append(cid) or FakeConn())
    monkeypatch.setattr(cmp, "resolve_key_candidates",
                        lambda source_id, tables: {})

    job_id = create_job("pg_compare", 11, {
        "source_connection_id": 11, "dest_connection_id": 12,
        "tables": [dict(t, in_src=True, in_dst=True) for t in _two_tables()],
        "targets": targets, "item_action": "COMPARE"})

    cmp.run_pg_compare_job(job_id)

    _untouched(job_id)
    assert opened == []


@pytest.mark.parametrize("targets", BAD_MAPS)
def test_diff_load_runner_rejects_a_bad_stored_map(monkeypatch, targets):
    opened = []
    monkeypatch.setattr(pdl, "open_pg", lambda cid, readonly=False:
                        opened.append(cid) or FakeConn())

    job_id = create_job("pg_diff_load", 11, {
        "source_connection_id": 11, "dest_connection_id": 12,
        "compare_job_id": None, "delete_missing": True,
        "tables": [dict(t, action="full", key_columns=[], in_dst=True)
                   for t in _two_tables()],
        "targets": targets, "expected": []})

    pdl.run_pg_diff_load_job(job_id)

    _untouched(job_id)
    assert opened == []


def test_runner_accepts_a_good_stored_map_in_short_form():
    items = [{"schema_name": "s", "table_name": "a"},
             {"schema_name": "s", "table_name": "b"}]

    assert st.validated_targets({"targets": {"s.a": "s.a2"}}, items) == {
        "s.a": "s.a2"}
    assert st.validated_targets({}, items) == {}