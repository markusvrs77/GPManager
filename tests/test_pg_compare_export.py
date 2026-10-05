# -*- coding: utf-8 -*-
"""
Выгрузка результатов сравнения баз PostgreSQL в Excel:
GET /api/pg/compare/<job_id>/export.xlsx?filter=...&q=...
"""

import io
import re

import pytest
from openpyxl import load_workbook

import modules.pg_compare as cmp
from job_manager import create_job
from modules import web_auth

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

HEADER = ["Схема", "Таблица", "Статус", "Ключ", "Источник ключа",
          "Строк в источнике", "Строк в приёмнике", "Добавить", "Изменить",
          "Удалить", "Всего отличий", "Диапазоны", "Сообщение",
          "Время сравнения", "Приёмник"]


@pytest.fixture
def compare_job():
    """Сравнение пары 81 → 82 со всеми группами статусов."""
    job_id = create_job("pg_compare", 81, {
        "source_connection_id": 81, "dest_connection_id": 82,
        "tables": [], "item_action": "COMPARE"})
    rows = [
        {"schema": "sales", "table": "orders", "status": "differs",
         "key_columns": ["id"], "key_source": "pk", "src_rows": 100,
         "dst_rows": 90, "to_insert": 10, "to_update": 5, "to_delete": 0,
         "chunked": {"checked": 64, "total": 64, "mismatched": 2}},
        {"schema": "sales", "table": "Items", "status": "no_dest",
         "src_rows": 7},
        {"schema": "sales", "table": "small", "status": "differs",
         "key_columns": ["id"], "key_source": "pk", "src_rows": 3,
         "dst_rows": 3, "to_insert": None, "to_update": 1, "to_delete": 0},
        {"schema": "hr", "table": "people", "status": "same",
         "key_columns": ["id"], "key_source": "pk", "src_rows": 5,
         "dst_rows": 5, "to_insert": 0, "to_update": 0, "to_delete": 0},
        {"schema": "hr", "table": "dups", "status": "duplicate_keys",
         "message": "дубликаты ключа в источнике"},
        {"schema": "hr", "table": "broken", "status": "error",
         "message": "relation does not exist"},
    ]
    for row in rows:
        cmp.save_result(job_id, row)
    column = {"name": "id", "collate_c": False}
    cmp.save_leaf(job_id, "sales", "orders", column,
                  {"lo": 1, "hi": 100, "depth": 1},
                  {"status": "differs", "src_rows": 50, "dst_rows": 45,
                   "to_insert": 5, "to_update": 2, "to_delete": 0})
    cmp.save_leaf(job_id, "sales", "orders", column,
                  {"lo": "ab", "depth": 2, "mode": "bucket"},
                  {"status": "differs", "src_rows": 50, "dst_rows": 45,
                   "to_insert": 5, "to_update": 3, "to_delete": 0})
    return job_id


def _book(response):
    assert response.status_code == 200, response.get_data(as_text=True)[:300]
    return load_workbook(io.BytesIO(response.get_data()))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]


def test_export_is_an_xlsx_with_results_sheet(client, compare_job):
    response = client.get("/api/pg/compare/%d/export.xlsx" % compare_job)

    assert response.headers["Content-Type"].startswith(XLSX)
    disposition = response.headers["Content-Disposition"]
    assert re.search(r"pg_compare_%d_\d{8}_\d{4}\.xlsx" % compare_job,
                     disposition), disposition

    book = _book(response)
    assert book.sheetnames[0] == "Результаты"
    ws = book["Результаты"]
    assert [c.value for c in ws[1]] == HEADER
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert ws["A1"].font.bold

    rows = _rows(ws)
    assert len(rows) == 6
    # по «всего отличий» по убыванию: 15, 1, затем нули / пустые
    assert rows[0][:3] == ["sales", "orders", "Отличается"]
    assert rows[0][3:11] == ["id", "PK", 100, 90, 10, 5, 0, 15]
    assert rows[0][11] == "64 / 64 / 2"
    assert rows[1][:2] == ["sales", "small"] and rows[1][10] == 1
    statuses = {r[1]: r[2] for r in rows}
    assert statuses["Items"] == "Нет в приёмнике"
    assert statuses["people"] == "Совпадает"
    assert statuses["dups"] == "Дубликаты ключа"
    assert {r[1]: r[12] for r in rows}["broken"] == "relation does not exist"


def test_needs_filter_keeps_only_tables_to_align(client, compare_job):
    book = _book(client.get("/api/pg/compare/%d/export.xlsx?filter=needs"
                            % compare_job))

    rows = _rows(book["Результаты"])
    assert sorted(r[1] for r in rows) == ["Items", "orders", "small"]
    assert {r[2] for r in rows} == {"Отличается", "Нет в приёмнике"}


def test_search_is_case_insensitive_over_schema_dot_table(client, compare_job):
    book = _book(client.get("/api/pg/compare/%d/export.xlsx?filter=all"
                            "&q=SALES.it" % compare_job))

    assert [r[:2] for r in _rows(book["Результаты"])] == [["sales", "Items"]]


def test_mismatched_ranges_sheet(client, compare_job):
    book = _book(client.get("/api/pg/compare/%d/export.xlsx" % compare_job))

    ws = book["Несовпавшие диапазоны"]
    assert [c.value for c in ws[1]] == [
        "Схема", "Таблица", "Колонка", "Режим", "От", "До", "NULL-диапазон",
        "Строк в источнике", "Строк в приёмнике", "Добавить", "Изменить",
        "Удалить"]
    assert _rows(ws) == [
        ["sales", "orders", "id", "диапазон", 1, 100, "нет", 50, 45, 5, 2, 0],
        ["sales", "orders", "id", "корзина", "ab", None, "нет", 50, 45, 5, 3,
         0],
    ]


def test_no_ranges_sheet_when_filter_hides_chunked_tables(client, compare_job):
    book = _book(client.get("/api/pg/compare/%d/export.xlsx?filter=same"
                            % compare_job))

    assert book.sheetnames == ["Результаты", "Сводка"]
    assert [r[1] for r in _rows(book["Результаты"])] == ["people"]


def test_summary_names_connections_without_passwords(client):
    from modules.connections import create_connection

    secret = "Pa55-very-secret"
    src = create_connection({"name": "prod-cash", "host": "h1",
                             "database_name": "cash", "username": "u",
                             "password": secret, "db_type": "postgres"})
    dst = create_connection({"name": "test-cash", "host": "h2",
                             "database_name": "cash", "username": "u",
                             "password": secret, "db_type": "postgres"})
    job_id = create_job("pg_compare", src, {
        "source_connection_id": src, "dest_connection_id": dst,
        "tables": [], "item_action": "COMPARE"})
    cmp.save_result(job_id, {"schema": "s", "table": "t", "status": "same"})
    cmp.save_result(job_id, {"schema": "s", "table": "u", "status": "error",
                             "message": "boom"})

    book = _book(client.get("/api/pg/compare/%d/export.xlsx" % job_id))

    summary = {r[0]: r[1] for r in _rows(book["Сводка"])}
    assert "prod-cash" in summary["Источник"]
    assert "test-cash" in summary["Приёмник"]
    assert summary["Таблиц всего"] == 2
    assert summary["Совпадает"] == 1 and summary["Ошибка"] == 1
    for ws in book.worksheets:
        for row in ws.iter_rows(values_only=True):
            assert all(secret not in str(v) for v in row if v is not None)


def test_export_of_foreign_or_missing_job_is_404(client):
    foreign = create_job("copy_pipe", 81, {"tables": []})

    assert client.get("/api/pg/compare/%d/export.xlsx"
                      % foreign).status_code == 404
    assert client.get("/api/pg/compare/987654/export.xlsx").status_code == 404


def test_unknown_filter_is_400(client, compare_job):
    response = client.get("/api/pg/compare/%d/export.xlsx?filter=nope"
                          % compare_job)

    assert response.status_code == 400


def test_export_needs_sync_view(as_user, compare_job):
    response = as_user("operator", overrides={"sync.view": False}).get(
        "/api/pg/compare/%d/export.xlsx" % compare_job)

    assert response.status_code == 403
    assert web_auth.POLICY["api_pg_compare_export"] == "sync.view"
